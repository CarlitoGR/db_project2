#!/usr/bin/env python3
"""Convert schedgen CSV output into Word (.docx) and Excel (.xlsx) documents.

Called from the Makefile ``docs`` target:

    python3 scripts/export_docs.py --input output \
        --docx output/Faculty_Schedules.docx --xlsx output/Faculty_Schedules.xlsx

The C program owns all scheduling logic and display formatting (e.g. "Monday & Wednesday
4:00 PM to 6:50 PM"). This script only lays those strings out in documents, so the text,
Word and Excel outputs never disagree.
"""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass, field
from pathlib import Path

NOTE_TEXT = "NOTE: Instructors are not to be interrupted during class times."
SESSION_TEXT = "Office hours apply to both sessions (7R1 and 7R2). * = see data note below the table."
HEADER_FILL = "D9D9D9"
ERROR_FILL = "F4CCCC"
WARN_FILL = "FFF2CC"


@dataclass
class Faculty:
    """One professor plus the rows that belong to them."""

    faculty_id: str
    name: str
    department: str
    location: str
    email: str
    phone: str
    sections: int
    enrollment: int
    required_hours: float
    assigned_hours: float
    classes: list[dict[str, str]] = field(default_factory=list)
    office_hours: list[dict[str, str]] = field(default_factory=list)


@dataclass
class Report:
    """Everything the documents need."""

    faculty: list[Faculty]
    unassigned: list[dict[str, str]]
    diagnostics: list[tuple[str, str]]


def read_csv(path: Path) -> list[dict[str, str]]:
    """Read a CSV file into a list of dicts, failing with a clear message if missing."""
    if not path.is_file():
        sys.exit(f"export_docs: missing {path} - run 'make run' first")
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def read_diagnostics(path: Path) -> list[tuple[str, str]]:
    """Parse 'LEVEL: message' lines from diagnostics.txt."""
    if not path.is_file():
        return []
    result = []
    for line in path.read_text(encoding="utf-8").splitlines():
        level, sep, message = line.partition(": ")
        if sep and level in {"ERROR", "WARNING", "INFO"}:
            result.append((level, message))
    return result


def load_report(input_dir: Path) -> Report:
    """Join faculty_summary.csv, classes.csv and office_hours.csv by faculty_id."""
    faculty = [
        Faculty(
            faculty_id=row["faculty_id"],
            name=row["name"],
            department=row["department"],
            location=row["office_hours_location"],
            email=row["email"],
            phone=row["phone"],
            sections=int(row["sections"]),
            enrollment=int(row["enrollment"]),
            required_hours=float(row["oh_required_hours"]),
            assigned_hours=float(row["oh_assigned_hours"]),
        )
        for row in read_csv(input_dir / "faculty_summary.csv")
    ]
    by_id = {fac.faculty_id: fac for fac in faculty}
    unassigned: list[dict[str, str]] = []
    for row in read_csv(input_dir / "classes.csv"):
        target = by_id.get(row["faculty_id"])
        (target.classes if target else unassigned).append(row)
    for row in read_csv(input_dir / "office_hours.csv"):
        if row["faculty_id"] in by_id:
            by_id[row["faculty_id"]].office_hours.append(row)
    return Report(faculty, unassigned, read_diagnostics(input_dir / "diagnostics.txt"))


def contact_line(fac: Faculty) -> str:
    """Single line of department / office-hours location / contact details."""
    parts = [fac.department, f"Office hours: {fac.location}", fac.email, fac.phone]
    return " | ".join(part for part in parts if part)


def class_rows(classes: list[dict[str, str]]) -> list[list[str]]:
    return [[c["course_display"], c["session"] or "-", c["times_display"], c["room"]] for c in classes]


def class_notes(classes: list[dict[str, str]]) -> list[str]:
    return [f"* {c['section_id']}: {c['note']}" for c in classes if c["note"] and c["meeting"] != "Asynchronous"]


# ------------------------------------------------------------------------------------------------
# Word
# ------------------------------------------------------------------------------------------------


def write_docx(report: Report, path: Path) -> None:
    """Data-issues page, then one page per professor: class schedule, office hours, note."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    def shade(cell: object, color: str) -> None:
        props = cell._tc.get_or_add_tcPr()  # noqa: SLF001 - python-docx has no public shading API
        fill = OxmlElement("w:shd")
        fill.set(qn("w:val"), "clear")
        fill.set(qn("w:color"), "auto")
        fill.set(qn("w:fill"), color)
        props.append(fill)

    def add_table(doc: object, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, text in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
            shade(cell, HEADER_FILL)
        for values in rows or [["(none)"] + [""] * (len(headers) - 1)]:
            cells = table.add_row().cells
            for idx, text in enumerate(values):
                cells[idx].text = ""
                run = cells[idx].paragraphs[0].add_run(text)
                run.bold = idx == 0
                run.font.size = Pt(10)
        table.autofit = False
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
            for row in table.rows:
                row.cells[idx].width = Inches(width)

    def add_heading(doc: object, text: str) -> None:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(12)
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(13)

    def add_small(doc: object, text: str, bold: bool = False) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(text)
        run.font.size = Pt(9)
        run.bold = bold

    class_widths = [2.6, 0.65, 2.95, 0.8]
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # Cover page: data issues the department should resolve.
    title = doc.add_heading("Faculty Teaching & Office-Hours Schedules", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total_sections = sum(f.sections for f in report.faculty) + len(report.unassigned)
    summary = doc.add_paragraph(f"{len(report.faculty)} faculty · {total_sections} sections")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "DATA ISSUES FOUND IN THE CLASS SCHEDULE")
    issues = [(lvl, msg) for lvl, msg in report.diagnostics if lvl in {"ERROR", "WARNING"}]
    add_table(doc, ["Severity", "Issue"], [[lvl, msg] for lvl, msg in issues], [0.9, 6.0])
    if report.unassigned:
        add_heading(doc, "SECTIONS WITH AN INSTRUCTOR NOT IN THE FACULTY FILE")
        add_table(doc, ["Course", "Session", "Times", "Location"], class_rows(report.unassigned), class_widths)

    for fac in report.faculty:
        doc.add_page_break()
        heading = doc.add_heading(f"{fac.name} ({fac.faculty_id})", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info = doc.add_paragraph(contact_line(fac))
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER

        add_heading(doc, "CLASS SCHEDULE")
        add_table(doc, ["Course", "Session", "Times", "Location"], class_rows(fac.classes), class_widths)
        for note in class_notes(fac.classes):
            add_small(doc, note)

        add_heading(doc, "OFFICE HOURS")
        add_table(
            doc,
            ["Day(s)", "Hours", "Location"],
            [[row["days_display"], row["hours_display"], row["location"]] for row in fac.office_hours],
            [3.0, 2.4, 1.5],
        )
        add_small(doc, f"Office hours per week: {fac.assigned_hours:g} h. {SESSION_TEXT}")
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_before = Pt(8)
        note.add_run(NOTE_TEXT).bold = True

    doc.save(str(path))


# ------------------------------------------------------------------------------------------------
# Excel
# ------------------------------------------------------------------------------------------------


def write_xlsx(report: Report, path: Path) -> None:
    """Summary, Classes, Office Hours, Data Issues, plus one sheet per professor."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)

    def write_block(ws: object, top: int, headers: list[str], rows: list[list[object]]) -> int:
        """Write a bordered table starting at row `top`; return the next free row."""
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=top, column=col, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        for r_offset, values in enumerate(rows, start=1):
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=top + r_offset, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="top")
        return top + len(rows) + 1

    def finish_table(ws: object, last_col: str, end_row: int) -> None:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{last_col}{max(end_row - 1, 1)}"
        widths: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(8, min(70, width + 2))

    def as_int(text: str) -> int | str:
        return int(text) if text.strip().lstrip("-").isdigit() else text

    all_classes = [c for fac in report.faculty for c in fac.classes] + report.unassigned
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    end = write_block(
        summary,
        1,
        ["Faculty ID", "Name", "Department", "Office-Hours Location", "Email", "Phone", "Sections", "Enrollment",
         "OH Required (h)", "OH Assigned (h)"],
        [
            [f.faculty_id, f.name, f.department, f.location, f.email, f.phone, f.sections, f.enrollment,
             f.required_hours, f.assigned_hours]
            for f in report.faculty
        ],
    )
    finish_table(summary, "J", end)

    classes = wb.create_sheet("Classes")
    end = write_block(
        classes,
        1,
        ["Section", "Title", "Session", "Instructor (as listed)", "Faculty", "Meeting", "Days", "Start", "End",
         "Room", "Credits", "Enrollment", "Note"],
        [
            [c["section_id"], c["title"], c["session"], c["instructor_as_listed"], c["faculty_name"], c["meeting"],
             c["days"], c["start_time"], c["end_time"], c["room"], as_int(c["credits"]), as_int(c["enrollment"]),
             c["note"]]
            for c in all_classes
        ],
    )
    finish_table(classes, "M", end)

    hours = wb.create_sheet("Office Hours")
    end = write_block(
        hours,
        1,
        ["Faculty ID", "Faculty", "Days", "Start", "End", "Location", "Arranged by"],
        [
            [r["faculty_id"], r["faculty_name"], r["days"], r["start_time"], r["end_time"], r["location"],
             r["source"]]
            for f in report.faculty
            for r in f.office_hours
        ],
    )
    finish_table(hours, "G", end)

    issues = wb.create_sheet("Data Issues")
    end = write_block(issues, 1, ["Severity", "Issue"], [[lvl, msg] for lvl, msg in report.diagnostics])
    for row in issues.iter_rows(min_row=2, max_row=end - 1):
        fill = {"ERROR": ERROR_FILL, "WARNING": WARN_FILL}.get(str(row[0].value))
        if fill:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=fill)
    finish_table(issues, "B", end)

    for fac in report.faculty:
        last_name = fac.name.split()[-1] if fac.name.split() else fac.faculty_id
        ws = wb.create_sheet(f"{fac.faculty_id} {last_name}"[:31])
        ws.cell(row=1, column=1, value=f"{fac.name} ({fac.faculty_id})").font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value=contact_line(fac))
        ws.cell(row=4, column=1, value="CLASS SCHEDULE").font = Font(bold=True, size=12)
        row = write_block(ws, 5, ["Course", "Session", "Times", "Location"], class_rows(fac.classes))
        for note in class_notes(fac.classes):
            ws.cell(row=row, column=1, value=note).font = Font(italic=True, size=9)
            row += 1
        ws.cell(row=row + 1, column=1, value="OFFICE HOURS").font = Font(bold=True, size=12)
        row = write_block(
            ws,
            row + 2,
            ["Day(s)", "Hours", "Location"],
            [[h["days_display"], h["hours_display"], h["location"]] for h in fac.office_hours],
        )
        ws.cell(row=row, column=1, value=SESSION_TEXT).font = Font(italic=True, size=9)
        ws.cell(row=row + 1, column=1, value=NOTE_TEXT).font = Font(bold=True)
        for col, width in zip("ABCD", (48, 20, 46, 18), strict=True):
            ws.column_dimensions[col].width = width

    wb.save(str(path))


# ------------------------------------------------------------------------------------------------
# Entry point
# ------------------------------------------------------------------------------------------------


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export schedgen CSV output to Word and Excel.")
    parser.add_argument("--input", type=Path, default=Path("output"), help="directory holding schedgen CSV files")
    parser.add_argument("--docx", type=Path, help="Word document to write")
    parser.add_argument("--xlsx", type=Path, help="Excel workbook to write")
    args = parser.parse_args(argv)
    if args.docx is None and args.xlsx is None:
        parser.error("specify --docx and/or --xlsx")
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    report = load_report(args.input)
    try:
        if args.docx is not None:
            write_docx(report, args.docx)
            print(f"export_docs: wrote {args.docx}")
        if args.xlsx is not None:
            write_xlsx(report, args.xlsx)
            print(f"export_docs: wrote {args.xlsx}")
    except ModuleNotFoundError as exc:
        print(
            f"export_docs: missing Python package '{exc.name}'. Install with:\n"
            "    python3 -m pip install -r requirements.txt   (or: make venv)",
            file=sys.stderr,
        )
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
